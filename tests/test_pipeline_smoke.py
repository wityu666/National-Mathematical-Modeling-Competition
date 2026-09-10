"""Synthetic script integration, not a mathematical or visual full-suite PASS."""

import hashlib
import json
import shutil
import subprocess
import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[1]
MANIFEST = ROOT / "cumcm-live-problem-analyst/scripts/build_problem_manifest.py"
COMPARE = ROOT / "cumcm-live-result-verifier/scripts/compare_runs.py"
LAYOUT = ROOT / "cumcm-live-layout-verifier/scripts/layout_preflight.py"
AUDIT = ROOT / "cumcm-live-final-auditor/scripts/audit_submission.py"

MODEL = """import csv
import json
import sys
from pathlib import Path

rows = list(csv.DictReader(Path(sys.argv[1]).open()))
x = [float(row["x"]) for row in rows]
y = [float(row["y"]) for row in rows]
n = len(x)
slope = (n * sum(a*b for a, b in zip(x, y)) - sum(x)*sum(y)) / (
    n * sum(a*a for a in x) - sum(x)**2
)
intercept = (sum(y) - slope*sum(x)) / n
output = Path(sys.argv[2])
output.mkdir()
result = {"slope": slope, "intercept": intercept, "prediction_x5": slope*5 + intercept}
(output / "summary.json").write_text(json.dumps(result, sort_keys=True))
(output / "metrics.csv").write_text("metric,value\\nprediction_x5," + str(result["prediction_x5"]) + "\\n")
"""


def command(script, *args, expected=0):
    completed = subprocess.run(
        [sys.executable, str(script), *map(str, args)],
        check=False, capture_output=True, text=True,
    )
    assert completed.returncode == expected, completed.stderr + completed.stdout
    return completed


def json_command(script, *args, expected=0):
    return json.loads(command(script, *args, expected=expected).stdout)


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def make_documents(tmp_path, model_source):
    reportlab = pytest.importorskip("reportlab")
    docx = pytest.importorskip("docx")
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.pdfgen import canvas

    font_path = Path(reportlab.__file__).parent / "fonts" / "Vera.ttf"
    pdfmetrics.registerFont(TTFont("SmokeVera", str(font_path)))
    pages = [
        ["Synthetic fixture abstract", "Known relation: y = 2*x + 1."],
        ["Problem and inputs", "Four observations: (1,3), (2,5), (3,7), (4,9)."],
        ["Model", "Least-squares slope = 2.0; intercept = 1.0."],
        ["Results", "Prediction at x=5: 11.0"],
        ["Independent check", "Endpoint slope: (9-3)/(4-1) = 2.0.",
         "Substitution gives intercept 1.0 and prediction 11.0."],
        ["Appendix: modeling source", *model_source.splitlines()],
    ]
    pdf = tmp_path / "paper.pdf"
    word = tmp_path / "paper.docx"
    document = docx.Document()
    drawing = canvas.Canvas(str(pdf))
    for index, lines in enumerate(pages):
        if index:
            document.add_page_break()
        drawing.setFont("SmokeVera", 10)
        for line_number, line in enumerate(lines):
            drawing.drawString(40, 790 - line_number * 16, line)
            document.add_paragraph(line)
        drawing.showPage()
    drawing.save()
    document.save(word)
    return pdf, word


def test_synthetic_pipeline_and_stale_evidence_signals(tmp_path):
    missing = [name for name in ("pdfinfo", "pdffonts", "pdftotext") if not shutil.which(name)]
    if missing:
        pytest.skip("Poppler required for actual PDF smoke test: " + ", ".join(missing))

    raw = tmp_path / "raw"
    raw.mkdir()
    data = raw / "data.csv"
    data.write_text("x,y\n1,3\n2,5\n3,7\n4,9\n")
    original_hash = digest(data)
    manifest = json_command(MANIFEST, raw, "--format", "json", "--sha256")
    record = next(row for row in manifest["files"] if row["relative_path"] == "data.csv")
    assert record["role"] == "data_attachment"
    assert record["sha256"] == original_hash

    model = tmp_path / "model.py"
    model.write_text(MODEL)
    initial, recheck = tmp_path / "initial", tmp_path / "recheck"
    command(model, data, initial)
    command(model, data, recheck)
    result = json.loads((initial / "summary.json").read_text())
    # Independently known answer from the synthetic relation and endpoint slope.
    assert result == {"slope": 2.0, "intercept": 1.0, "prediction_x5": 11.0}
    comparison = json_command(COMPARE, initial, recheck, "--json")
    assert comparison["status"] == "PASS"
    assert comparison["comparison_mode"] == "exact"
    assert digest(data) == original_hash  # producer did not modify inputs

    pdf, word = make_documents(tmp_path, MODEL)
    from docx import Document
    word_text = "\n".join(p.text for p in Document(word).paragraphs)
    extracted = subprocess.run(
        ["pdftotext", str(pdf), "-"], check=True, capture_output=True, text=True,
    ).stdout
    assert "Prediction at x=5: 11.0" in word_text
    assert "Prediction at x=5: 11.0" in extracted

    policy_path = tmp_path / "team-policy.json"
    policy = {
        "schema_version": 1, "policy_id": "synthetic-smoke",
        "rationale": "Four-page body for a synthetic integration fixture.",
        "min_main_pages": 4, "max_main_pages": 4,
    }
    policy_path.write_text(json.dumps(policy))
    layout_args = (
        pdf, "--source", word, "--team-policy", policy_path,
        "--abstract-end-page", "1", "--main-start-page", "2",
        "--appendix-start-page", "6", "--appendix-code-page", "6", "--json",
    )
    layout = json_command(LAYOUT, *layout_args)
    assert layout["status"] == "PRECHECK_PASS"
    assert layout["page_limit"]["main_body_pages"] == 4
    assert layout["page_policy"]["sha256"] == digest(policy_path)
    assert layout["pdf_sha256"] == digest(pdf)
    assert layout["visual_qa_required"] is True
    assert layout["appendix_code"]["status"] == "DECLARED"  # still needs visual confirmation

    submission = tmp_path / "submission"
    submission.mkdir()
    shutil.copy2(pdf, submission / "paper.pdf")
    shutil.copy2(model, submission / "model.py")
    audit = json_command(AUDIT, submission, "--json")
    assert audit["status"] == "PASS"  # directory checks only
    assert {f["path"] for f in audit["files"]} == {"paper.pdf", "model.py"}
    audited_pdf = next(f for f in audit["files"] if f["path"] == "paper.pdf")
    assert audited_pdf["sha256"] == layout["pdf_sha256"]

    # The old tolerances would have hidden this difference.
    drifted = dict(result, prediction_x5=11.0000000005)
    (recheck / "summary.json").write_text(json.dumps(drifted))
    assert json_command(COMPARE, initial, recheck, "--json", expected=1)["status"] == "BLOCKED"

    # New policy requires a fresh page judgment and cannot inherit the old report.
    policy.update(min_main_pages=5, max_main_pages=5)
    policy_path.write_text(json.dumps(policy))
    new_layout = json_command(LAYOUT, *layout_args, expected=1)
    assert new_layout["page_policy"]["sha256"] != layout["page_policy"]["sha256"]
    assert any(i["code"] == "main-body-under-page-floor" for i in new_layout["issues"])

    # Changed input is visible in both provenance and the newly produced results.
    data.write_text("x,y\n1,3\n2,5\n3,7\n4,10\n")
    changed_manifest = json_command(MANIFEST, raw, "--format", "json", "--sha256")
    assert changed_manifest["files"][0]["sha256"] != record["sha256"]
    new_run = tmp_path / "changed-input-run"
    command(model, data, new_run)
    assert json_command(COMPARE, initial, new_run, "--json", expected=1)["status"] == "BLOCKED"
